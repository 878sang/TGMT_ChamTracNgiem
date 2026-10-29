"""
Module: excel_answer_handler.py
Phụ trách: Thành viên 3 (Backend / Data)
Chức năng: 
  - Xuất bộ đáp án đúng ra file Excel (.xlsx) hoặc file Text (.txt)
  - Đọc (Import) bộ đáp án đúng từ file Excel
  - Hỗ trợ chuyển đổi chuỗi nhập nhanh (vd: "A B C D A B...") thành dictionary
"""

import os
import sys
from typing import Dict, Union, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass



def parse_quick_answer_string(raw_text: str) -> Dict[int, str]:
    """
    Chuyển đổi chuỗi nhập nhanh thành dictionary câu hỏi và đáp án.
    Ví dụ: 'A B C D B A' hoặc '1A 2B 3C 4D' hoặc 'ABCD...'
    Trả về: {1: 'A', 2: 'B', 3: 'C', ...}
    """
    # Nếu là chuỗi phân tách bởi dấu cách, phẩy hoặc tab
    cleaned = raw_text.replace(',', ' ').replace(';', ' ').replace('\n', ' ').strip()
    tokens = [t.strip().upper() for t in cleaned.split() if t.strip()]
    
    result = {}
    valid_options = {'A', 'B', 'C', 'D', 'E', 'F'}
    
    # Trường hợp 1: Giáo viên nhập "A B C D A B..."
    if all(len(t) == 1 and t in valid_options for t in tokens):
        for idx, ans in enumerate(tokens, start=1):
            result[idx] = ans
        return result

    # Trường hợp 2: Giáo viên nhập "1A 2B 3C..." hoặc "1.A 2.B"
    idx_counter = 1
    for token in tokens:
        # Tách số và chữ nếu dính liền
        digits = "".join([c for c in token if c.isdigit()])
        alphas = "".join([c for c in token if c.isalpha()]).upper()
        
        if alphas in valid_options:
            q_num = int(digits) if digits else idx_counter
            result[q_num] = alphas
            idx_counter = q_num + 1

    # Trường hợp 3: Nhập dính liền một chuỗi chữ cái "ABCDABCD..."
    if not result and len(tokens) == 1 and len(tokens[0]) > 1:
        single_str = tokens[0].upper()
        if all(c in valid_options for c in single_str):
            for idx, c in enumerate(single_str, start=1):
                result[idx] = c

    return result


def export_answer_key_to_excel(
    answers: Dict[Union[int, str], str], 
    output_filepath: str, 
    exam_title: str = "BÀI KIỂM TRA TRẮC NGHIỆM",
    subject: str = "",
    exam_class: str = ""
) -> bool:
    """
    Xuất bộ đáp án đúng ra file Excel với định dạng chuyên nghiệp.
    - answers: Dictionary dạng {1: 'A', 2: 'B', 3: 'C', ...}
    - output_filepath: Đường dẫn file .xlsx muốn lưu
    """
    try:
        os.makedirs(os.path.dirname(os.path.abspath(output_filepath)), exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Dap_An_Chuan"

        # Tiêu đề chính
        ws.merge_cells("A1:B1")
        ws["A1"] = exam_title.upper()
        ws["A1"].font = Font(name="Arial", size=13, bold=True, color="FFFFFF")
        ws["A1"].fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 30

        # Thông tin môn học & lớp
        sub_info = f"Môn học: {subject}" if subject else ""
        if exam_class:
            sub_info += f" | Lớp: {exam_class}"
        if not sub_info:
            sub_info = f"Tổng số câu: {len(answers)}"

        ws.merge_cells("A2:B2")
        ws["A2"] = sub_info
        ws["A2"].font = Font(name="Arial", size=10, italic=True)
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 20

        # Header bảng
        headers = ["Câu", "Đáp án đúng"]
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col_idx, value=h)
            cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[3].height = 24

        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9")
        )

        # Ghi các câu hỏi theo thứ tự số nguyên
        sorted_keys = sorted(answers.keys(), key=lambda x: int(x) if str(x).isdigit() else str(x))
        for row_idx, q_num in enumerate(sorted_keys, start=4):
            c1 = ws.cell(row=row_idx, column=1, value=int(q_num) if str(q_num).isdigit() else q_num)
            c2 = ws.cell(row=row_idx, column=2, value=str(answers[q_num]).strip().upper())
            
            c1.alignment = Alignment(horizontal="center", vertical="center")
            c2.alignment = Alignment(horizontal="center", vertical="center")
            c2.font = Font(name="Arial", size=11, bold=True, color="002060")

            c1.border = thin_border
            c2.border = thin_border

            # Tô màu xen kẽ
            if (row_idx - 3) % 2 == 0:
                c1.fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
                c2.fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

        # Canh chỉnh độ rộng cột
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 18

        wb.save(output_filepath)
        return True
    except Exception as e:
        print(f"Lỗi khi xuất file Excel đáp án: {e}")
        return False


def import_answer_key_from_excel(filepath: str, sheet_name: Optional[str] = None) -> Dict[int, str]:
    """
    Đọc đáp án từ file Excel (.xlsx).
    Tự động tìm dòng bắt đầu có chữ 'Câu' hoặc số 1.
    Trả về: {1: 'A', 2: 'B', 3: 'C', ...}
    """
    if not os.path.exists(filepath):
        raise FileNotFoundError(f"Không tìm thấy file: {filepath}")

    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active

    answers = {}
    valid_options = {"A", "B", "C", "D", "E", "F"}

    for row in ws.iter_rows(values_only=True):
        if not row or len(row) < 2:
            continue
        col1, col2 = row[0], row[1]
        
        # Bỏ qua nếu giá trị rỗng
        if col1 is None or col2 is None:
            continue
        
        # Kiểm tra xem col1 có phải là số câu hỏi hay không
        try:
            q_num = int(col1)
            ans = str(col2).strip().upper()
            if ans in valid_options:
                answers[q_num] = ans
        except (ValueError, TypeError):
            # Không phải số (có thể là dòng tiêu đề như 'Câu' | 'Đáp án') -> bỏ qua
            continue

    return answers


def export_answer_key_to_txt(answers: Dict[Union[int, str], str], output_filepath: str) -> bool:
    """
    Xuất đáp án dạng chuỗi text đơn giản để lưu trữ hoặc copy nhanh.
    Ví dụ:
    1: A
    2: B
    ...
    Chuỗi nhanh: A B C D ...
    """
    try:
        os.makedirs(os.path.dirname(os.path.abspath(output_filepath)), exist_ok=True)
        sorted_keys = sorted(answers.keys(), key=lambda x: int(x) if str(x).isdigit() else str(x))
        quick_str = " ".join([str(answers[k]).upper() for k in sorted_keys])
        
        with open(output_filepath, "w", encoding="utf-8") as f:
            f.write("# ĐÁP ÁN BÀI KIỂM TRA\n")
            f.write(f"# DẠNG CHUỖI NHANH: {quick_str}\n\n")
            for k in sorted_keys:
                f.write(f"Câu {k}: {str(answers[k]).upper()}\n")
        return True
    except Exception as e:
        print(f"Lỗi khi xuất file TXT đáp án: {e}")
        return False


def export_thpt_answer_key_to_excel(
    exam_code: str,
    part1_answers: Dict[int, str],
    part2_answers: Dict[int, Dict[str, str]],
    part3_answers: Dict[int, str],
    output_filepath: str,
    subject: str = "Toán"
) -> bool:
    """
    Xuất bộ đáp án chuẩn THPT mới (Phần I, II, III) ra file Excel.
    - part1_answers: {1: 'A', 2: 'B', ...} (12 câu)
    - part2_answers: {1: {'a': 'Đ', 'b': 'S', 'c': 'Đ', 'd': 'Đ'}, ...} (4 câu)
    - part3_answers: {1: '-2.5', 2: '15', ...} (6 câu)
    """
    try:
        os.makedirs(os.path.dirname(os.path.abspath(output_filepath)), exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Ma_De_{exam_code}"

        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        sec_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
        stripe_fill = PatternFill(start_color="F2F4F7", end_color="F2F4F7", fill_type="solid")
        white_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        bold_font = Font(name="Arial", size=11, bold=True)
        center_al = Alignment(horizontal="center", vertical="center")

        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"), right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"), bottom=Side(style="thin", color="D9D9D9")
        )

        # Header chung
        ws.merge_cells("A1:D1")
        ws["A1"] = f"ĐÁP ÁN BÀI THI THPT MÔN {subject.upper()} - MÃ ĐỀ {exam_code}"
        ws["A1"].font = Font(name="Arial", size=13, bold=True, color="FFFFFF")
        ws["A1"].fill = header_fill
        ws["A1"].alignment = center_al

        # PHẦN I
        curr_row = 3
        ws.merge_cells(f"A{curr_row}:D{curr_row}")
        ws[f"A{curr_row}"] = "PHẦN I: TRẮC NGHIỆM 4 LỰA CHỌN (3.0 ĐIỂM)"
        ws[f"A{curr_row}"].font = white_font
        ws[f"A{curr_row}"].fill = sec_fill
        ws[f"A{curr_row}"].alignment = center_al

        curr_row += 1
        ws[f"A{curr_row}"] = "Câu"
        ws[f"B{curr_row}"] = "Đáp án"
        ws[f"A{curr_row}"].font = bold_font; ws[f"A{curr_row}"].alignment = center_al
        ws[f"B{curr_row}"].font = bold_font; ws[f"B{curr_row}"].alignment = center_al

        for q_num in sorted(part1_answers.keys(), key=lambda x: int(x)):
            curr_row += 1
            c1 = ws.cell(row=curr_row, column=1, value=f"Câu {q_num}")
            c2 = ws.cell(row=curr_row, column=2, value=str(part1_answers[q_num]).upper())
            c1.alignment = center_al; c2.alignment = center_al
            c1.border = thin_border; c2.border = thin_border
            if q_num % 2 == 0:
                c1.fill = stripe_fill; c2.fill = stripe_fill

        # PHẦN II
        curr_row += 2
        ws.merge_cells(f"A{curr_row}:D{curr_row}")
        ws[f"A{curr_row}"] = "PHẦN II: TRẮC NGHIỆM ĐÚNG / SAI (4.0 ĐIỂM)"
        ws[f"A{curr_row}"].font = white_font
        ws[f"A{curr_row}"].fill = sec_fill
        ws[f"A{curr_row}"].alignment = center_al

        curr_row += 1
        ws[f"A{curr_row}"] = "Câu"
        ws[f"B{curr_row}"] = "Ý (Lệnh hỏi)"
        ws[f"C{curr_row}"] = "Đáp án (Đ/S)"
        for col_name in [f"A{curr_row}", f"B{curr_row}", f"C{curr_row}"]:
            ws[col_name].font = bold_font; ws[col_name].alignment = center_al

        for q_num in sorted(part2_answers.keys(), key=lambda x: int(x)):
            sub_dict = part2_answers[q_num]
            for sub_key in sorted(sub_dict.keys()):
                curr_row += 1
                c1 = ws.cell(row=curr_row, column=1, value=f"Câu {q_num}")
                c2 = ws.cell(row=curr_row, column=2, value=f"Ý {sub_key}")
                c3 = ws.cell(row=curr_row, column=3, value=str(sub_dict[sub_key]).upper())
                for cell in [c1, c2, c3]:
                    cell.alignment = center_al
                    cell.border = thin_border
                if q_num % 2 == 0:
                    c1.fill = stripe_fill; c2.fill = stripe_fill; c3.fill = stripe_fill

        # PHẦN III
        curr_row += 2
        ws.merge_cells(f"A{curr_row}:D{curr_row}")
        ws[f"A{curr_row}"] = "PHẦN III: TRẮC NGHIỆM TRẢ LỜI NGẮN (3.0 ĐIỂM)"
        ws[f"A{curr_row}"].font = white_font
        ws[f"A{curr_row}"].fill = sec_fill
        ws[f"A{curr_row}"].alignment = center_al

        curr_row += 1
        ws[f"A{curr_row}"] = "Câu"
        ws[f"B{curr_row}"] = "Kết quả số"
        ws[f"A{curr_row}"].font = bold_font; ws[f"A{curr_row}"].alignment = center_al
        ws[f"B{curr_row}"].font = bold_font; ws[f"B{curr_row}"].alignment = center_al

        for q_num in sorted(part3_answers.keys(), key=lambda x: int(x)):
            curr_row += 1
            c1 = ws.cell(row=curr_row, column=1, value=f"Câu {q_num}")
            c2 = ws.cell(row=curr_row, column=2, value=str(part3_answers[q_num]))
            c1.alignment = center_al; c2.alignment = center_al
            c1.border = thin_border; c2.border = thin_border
            if q_num % 2 == 0:
                c1.fill = stripe_fill; c2.fill = stripe_fill

        ws.column_dimensions["A"].width = 16
        ws.column_dimensions["B"].width = 16
        ws.column_dimensions["C"].width = 16
        ws.column_dimensions["D"].width = 16

        wb.save(output_filepath)
        return True
    except Exception as e:
        print(f"Lỗi khi xuất file Excel THPT: {e}")
        return False


def import_thpt_answer_key_from_excel(filepath: str, sheet_name: Optional[str] = None) -> dict:
    """
    Đọc đáp án đề thi THPT 3 phần từ file Excel.
    Trả về:
    {
        "part1": {1: 'A', 2: 'B', ...},
        "part2": {1: {'a': 'Đ', 'b': 'S', ...}, ...},
        "part3": {1: '-2.5', 2: '15', ...}
    }
    """
    if not os.path.exists(filepath):
        raise FileNotFoundError(f"Không tìm thấy file: {filepath}")

    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active

    result = {
        "part1": {},
        "part2": {},
        "part3": {}
    }

    current_section = None

    for row in ws.iter_rows(values_only=True):
        if not row or not any(row):
            continue

        first_val = str(row[0]).strip() if row[0] is not None else ""

        # Phát hiện phần (kiểm tra III và II trước I)
        val_upper = first_val.upper()
        if "PHẦN III" in val_upper or "PHAN III" in val_upper:
            current_section = 3
            continue
        elif "PHẦN II" in val_upper or "PHAN II" in val_upper:
            current_section = 2
            continue
        elif "PHẦN I" in val_upper or "PHAN I" in val_upper:
            current_section = 1
            continue


        # Đọc dữ liệu theo từng phần
        if current_section == 1:
            # Dòng: Câu 1 | A
            if "CÂU" in first_val.upper() and len(row) >= 2 and row[1]:
                q_num = "".join([c for c in first_val if c.isdigit()])
                if q_num:
                    ans = str(row[1]).strip().upper()
                    if ans in {"A", "B", "C", "D"}:
                        result["part1"][int(q_num)] = ans

        elif current_section == 2:
            # Dòng: Câu 1 | Ý a | Đ
            if "CÂU" in first_val.upper() and len(row) >= 3 and row[1] and row[2]:
                q_num = "".join([c for c in first_val if c.isdigit()])
                sub_token = str(row[1]).strip().lower()
                sub_key = "".join([c for c in sub_token if c in {'a', 'b', 'c', 'd'}])
                ans = str(row[2]).strip().upper()
                if q_num and sub_key and ans in {"Đ", "S", "D", "T", "F"}:
                    ans_standard = "Đ" if ans in {"Đ", "D", "T"} else "S"
                    q_int = int(q_num)
                    if q_int not in result["part2"]:
                        result["part2"][q_int] = {}
                    result["part2"][q_int][sub_key] = ans_standard

        elif current_section == 3:
            # Dòng: Câu 1 | 15
            if "CÂU" in first_val.upper() and len(row) >= 2 and row[1] is not None:
                q_num = "".join([c for c in first_val if c.isdigit()])
                if q_num:
                    ans_val = str(row[1]).strip().replace(",", ".")
                    result["part3"][int(q_num)] = ans_val

    return result


# ------------------------------------------------------------------------------
# Chạy thử nghiệm độc lập module
# ------------------------------------------------------------------------------
if __name__ == "__main__":
    PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    thpt_file = os.path.join(PROJECT_ROOT, "templates", "mau_dap_an_thpt_2025.xlsx")
    if os.path.exists(thpt_file):
        imported_thpt = import_thpt_answer_key_from_excel(thpt_file)
        print("Test import THPT 2025 thanh cong:")
        print("Part 1 (4 lua chon):", imported_thpt["part1"])
        print("Part 2 (Dung/Sai):", imported_thpt["part2"])
        print("Part 3 (Tra loi ngan):", imported_thpt["part3"])
    else:
        print(f"File mau khong ton tai tai: {thpt_file}")


