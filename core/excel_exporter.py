"""
Module: excel_exporter.py
Xuất báo cáo kết quả thi linh hoạt theo số câu (20, 40, 50 câu) ra file Excel:
  - Sheet 1: Bảng điểm tổng hợp (Mã HS, Tên file, Đúng, Sai, Trống, Điểm, Xếp loại, Trạng thái)
  - Sheet 2: Chi tiết các câu bài làm của từng học sinh
  - Sheet 3: Thống kê phổ điểm & Tỷ lệ đúng/sai theo từng câu hỏi
"""

import os
from typing import List, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from core.models import Submission


def export_exam_results_to_excel(
    submissions: List[Submission], 
    answer_key: Dict[int, str], 
    output_filepath: str,
    exam_title: str = "BẢNG ĐIỂM BÀI THI TRẮC NGHIỆM"
) -> bool:
    """Xuất danh sách kết quả học sinh ra file Excel với 3 Sheet chuyên nghiệp"""
    try:
        os.makedirs(os.path.dirname(os.path.abspath(output_filepath)), exist_ok=True)
        wb = openpyxl.Workbook()
        total_q = len(answer_key) or 50

        fill_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        fill_subheader = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
        fill_stripe = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
        fill_green = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        fill_red = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
        fill_yellow = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

        font_title = Font(name="Arial", size=13, bold=True, color="FFFFFF")
        font_header = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        font_bold = Font(name="Arial", size=10, bold=True)
        font_regular = Font(name="Arial", size=10)

        al_center = Alignment(horizontal="center", vertical="center")
        al_left = Alignment(horizontal="left", vertical="center")

        thin_side = Side(style="thin", color="D9D9D9")
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        # -------------------------------------------------------------
        # SHEET 1: BẢNG ĐIỂM TỔNG HỢP
        # -------------------------------------------------------------
        ws1 = wb.active
        ws1.title = "Bang_Diem_Tong_Hop"
        ws1.views.sheetView[0].showGridLines = True

        ws1.merge_cells("A1:J1")
        cell_title = ws1["A1"]
        cell_title.value = f"{exam_title.upper()} ({total_q} CÂU)"
        cell_title.font = font_title
        cell_title.fill = fill_header
        cell_title.alignment = al_center
        ws1.row_dimensions[1].height = 36

        headers_s1 = [
            "STT", "Mã học sinh", "Tên file ảnh", "Mã đề", 
            f"Số câu đúng (/{total_q})", "Số câu sai", "Bỏ trống", "Điểm số", "Xếp loại", "Trạng thái / Ghi chú"
        ]

        for col_idx, h_text in enumerate(headers_s1, start=1):
            cell = ws1.cell(row=2, column=col_idx, value=h_text)
            cell.font = font_header
            cell.fill = fill_subheader
            cell.alignment = al_center
            cell.border = thin_border
        ws1.row_dimensions[2].height = 26

        for row_idx, sub in enumerate(submissions, start=3):
            is_stripe = (row_idx % 2 == 0)
            row_fill = fill_stripe if is_stripe else None

            stt = row_idx - 2
            st_id = sub.student_id or sub.submission_id
            fname = os.path.basename(sub.image_path) if sub.image_path else ""
            
            if sub.score >= 8.0:
                rank = "Giỏi"
            elif sub.score >= 6.5:
                rank = "Khá"
            elif sub.score >= 5.0:
                rank = "Trung bình"
            else:
                rank = "Yếu"

            note = sub.error_message if sub.status == "Lỗi" else sub.warning_message
            if not note and sub.status == "Thành công":
                note = "Hợp lệ"

            row_data = [
                stt, st_id, fname, sub.exam_code,
                sub.correct_count, sub.wrong_count, sub.blank_count,
                sub.score if sub.status != "Lỗi" else 0.0,
                rank if sub.status != "Lỗi" else "Lỗi",
                note
            ]

            for col_idx, val in enumerate(row_data, start=1):
                cell = ws1.cell(row=row_idx, column=col_idx, value=val)
                cell.font = font_regular
                cell.border = thin_border
                cell.alignment = al_center if col_idx not in [2, 3, 10] else al_left
                if row_fill:
                    cell.fill = row_fill

                if col_idx == 8 and sub.status != "Lỗi":
                    cell.font = font_bold
                    if sub.score >= 8.0:
                        cell.fill = fill_green
                    elif sub.score < 5.0:
                        cell.fill = fill_red

                if col_idx == 10 and sub.status == "Lỗi":
                    cell.fill = fill_red

            ws1.row_dimensions[row_idx].height = 22

        for col in ws1.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws1.column_dimensions[col_letter].width = max(max_len + 3, 11)
        ws1.column_dimensions["C"].width = 28
        ws1.column_dimensions["J"].width = 35

        # -------------------------------------------------------------
        # SHEET 2: CHI TIẾT BÀI LÀM
        # -------------------------------------------------------------
        ws2 = wb.create_sheet(title="Chi_Tiet_Bai_Lam")
        ws2.views.sheetView[0].showGridLines = True

        headers_s2 = ["STT", "Mã HS", "Tên file"] + [f"C{i}" for i in range(1, total_q + 1)]
        for col_idx, h_text in enumerate(headers_s2, start=1):
            cell = ws2.cell(row=1, column=col_idx, value=h_text)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = al_center
            cell.border = thin_border
        ws2.row_dimensions[1].height = 26

        # Hàng KEY
        ws2.cell(row=2, column=1, value="KEY").font = font_bold
        ws2.cell(row=2, column=2, value="ĐÁP ÁN ĐỀ").font = font_bold
        ws2.cell(row=2, column=3, value="CHÍNH THỨC").font = font_bold
        for c in range(1, 4):
            ws2.cell(row=2, column=c).fill = fill_yellow
            ws2.cell(row=2, column=c).alignment = al_center
            ws2.cell(row=2, column=c).border = thin_border

        for q in range(1, total_q + 1):
            cell = ws2.cell(row=2, column=3 + q, value=answer_key.get(q, ""))
            cell.font = font_bold
            cell.fill = fill_yellow
            cell.alignment = al_center
            cell.border = thin_border
        ws2.row_dimensions[2].height = 22

        # Từng học sinh
        for row_idx, sub in enumerate(submissions, start=3):
            stt = row_idx - 2
            st_id = sub.student_id or sub.submission_id
            fname = os.path.basename(sub.image_path) if sub.image_path else ""

            ws2.cell(row=row_idx, column=1, value=stt).alignment = al_center
            ws2.cell(row=row_idx, column=2, value=st_id).alignment = al_left
            ws2.cell(row=row_idx, column=3, value=fname).alignment = al_left
            for c in range(1, 4):
                ws2.cell(row=row_idx, column=c).font = font_regular
                ws2.cell(row=row_idx, column=c).border = thin_border

            for q in range(1, total_q + 1):
                ans = sub.answers.get(q)
                key = answer_key.get(q, "")
                disp_val = ""
                if ans is None:
                    disp_val = "-"
                elif isinstance(ans, (list, tuple)):
                    disp_val = ",".join(ans)
                else:
                    disp_val = str(ans)

                cell = ws2.cell(row=row_idx, column=3 + q, value=disp_val)
                cell.font = font_regular
                cell.alignment = al_center
                cell.border = thin_border

                if sub.status != "Lỗi":
                    if disp_val == key and key:
                        cell.fill = fill_green
                    else:
                        cell.fill = fill_red

            ws2.row_dimensions[row_idx].height = 20

        ws2.column_dimensions["A"].width = 6
        ws2.column_dimensions["B"].width = 14
        ws2.column_dimensions["C"].width = 24
        for q in range(1, total_q + 1):
            col_letter = openpyxl.utils.get_column_letter(3 + q)
            ws2.column_dimensions[col_letter].width = 5.5

        # -------------------------------------------------------------
        # SHEET 3: THỐNG KÊ PHỔ ĐIỂM
        # -------------------------------------------------------------
        ws3 = wb.create_sheet(title="Thong_Ke_Pho_Diem")
        ws3.views.sheetView[0].showGridLines = True

        valid_subs = [s for s in submissions if s.status != "Lỗi"]
        total_valid = len(valid_subs)

        ws3["A1"] = f"BÁO CÁO THỐNG KÊ PHỔ ĐIỂM ({total_q} CÂU)"
        ws3["A1"].font = font_title
        ws3["A1"].fill = fill_header
        ws3.merge_cells("A1:D1")
        ws3.row_dimensions[1].height = 32

        scores = [s.score for s in valid_subs]
        avg_score = round(sum(scores) / max(total_valid, 1), 2) if scores else 0.0
        max_score = max(scores) if scores else 0.0
        min_score = min(scores) if scores else 0.0

        stats_summary = [
            ("Tổng số bài xử lý:", len(submissions)),
            ("Số bài hợp lệ:", total_valid),
            ("Số bài lỗi:", len(submissions) - total_valid),
            ("Điểm trung bình:", avg_score),
            ("Điểm cao nhất:", max_score),
            ("Điểm thấp nhất:", min_score),
        ]

        for r_i, (k, v) in enumerate(stats_summary, start=3):
            c_k = ws3.cell(row=r_i, column=1, value=k)
            c_k.font = font_bold
            c_v = ws3.cell(row=r_i, column=2, value=v)
            c_v.font = font_bold
            c_v.alignment = al_center

        ws3["A11"] = "TỶ LỆ LÀM ĐÚNG THEO TỪNG CÂU HỎI"
        ws3["A11"].font = font_header
        ws3["A11"].fill = fill_subheader
        ws3.merge_cells("A11:D11")

        headers_s3 = ["Câu hỏi", "Đáp án đề", "Số học sinh đúng", "Tỷ lệ đúng (%)"]
        for c_i, h in enumerate(headers_s3, start=1):
            cell = ws3.cell(row=12, column=c_i, value=h)
            cell.font = font_header
            cell.fill = fill_subheader
            cell.alignment = al_center
            cell.border = thin_border

        for q in range(1, total_q + 1):
            key = answer_key.get(q, "")
            correct_cnt = sum(1 for s in valid_subs if s.answers.get(q) == key)
            pct = round((correct_cnt / max(total_valid, 1)) * 100, 1)

            r = 12 + q
            ws3.cell(row=r, column=1, value=f"Câu {q:02d}").alignment = al_center
            ws3.cell(row=r, column=2, value=key).alignment = al_center
            ws3.cell(row=r, column=3, value=correct_cnt).alignment = al_center
            ws3.cell(row=r, column=4, value=f"{pct}%").alignment = al_center

            for c in range(1, 5):
                ws3.cell(row=r, column=c).border = thin_border
                ws3.cell(row=r, column=c).font = font_regular

        ws3.column_dimensions["A"].width = 16
        ws3.column_dimensions["B"].width = 14
        ws3.column_dimensions["C"].width = 18
        ws3.column_dimensions["D"].width = 18

        wb.save(output_filepath)
        return True

    except Exception as e:
        print(f"Lỗi xuất Excel: {e}")
        return False
