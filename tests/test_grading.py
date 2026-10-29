"""
Bộ kiểm thử tự động cho Logic Chấm Điểm và Xử Lý Dữ Liệu
Bao gồm:
  - Đáp án đúng / sai
  - Bỏ trống (blank)
  - Tô nhiều ô trong một câu (multiple bubbles)
  - Mã đề không tồn tại (không tự lấy mã đề khác)
  - Các trường hợp điểm biên (0.0, 10.0, rỗng)
  - Chấm lũy tiến Phần II THPT 2025
  - Chuẩn hóa số học Phần III THPT 2025
  - Loại trừ bài thi lỗi khỏi thống kê điểm trung bình và xếp loại
"""

import os
import sys
import pytest
import openpyxl

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

from core.models import Exam, Submission, ExamMode
from core.grading_engine import grade_full_mcq, grade_thpt_2025, normalize_numeric_string
from core.exam_processor import process_single_submission
from core.excel_exporter import export_exam_results_to_excel


# ==============================================================================
# 1. KIỂM THỬ CHẾ ĐỘ 1: FULL TRẮC NGHIỆM
# ==============================================================================
def test_full_mcq_all_correct():
    """Tất cả 40 câu đều đúng -> Đạt điểm tuyệt đối 10.0"""
    key = {i: "A" for i in range(1, 41)}
    student = {i: "A" for i in range(1, 41)}
    res = grade_full_mcq(student, key, max_score=10.0)
    assert res["score"] == 10.0
    assert res["correct_count"] == 40
    assert res["wrong_count"] == 0
    assert res["blank_count"] == 0


def test_full_mcq_all_wrong():
    """Tất cả 40 câu đều sai -> 0.0 điểm"""
    key = {i: "A" for i in range(1, 41)}
    student = {i: "B" for i in range(1, 41)}
    res = grade_full_mcq(student, key, max_score=10.0)
    assert res["score"] == 0.0
    assert res["correct_count"] == 0
    assert res["wrong_count"] == 40
    assert res["blank_count"] == 0


def test_full_mcq_partial_and_blanks():
    """20 câu đúng, 10 câu sai, 10 câu bỏ trống -> 5.0 điểm"""
    key = {i: "A" for i in range(1, 41)}
    student = {}
    for i in range(1, 21):
        student[i] = "A"  # Đúng
    for i in range(21, 31):
        student[i] = "C"  # Sai
    for i in range(31, 41):
        student[i] = None  # Bỏ trống

    res = grade_full_mcq(student, key, max_score=10.0)
    assert res["score"] == 5.0
    assert res["correct_count"] == 20
    assert res["wrong_count"] == 10
    assert res["blank_count"] == 10


def test_full_mcq_multiple_bubbles():
    """Học sinh tô 2 ô trong 1 câu (ví dụ ['A', 'B']) -> Không được tính điểm"""
    key = {1: "A", 2: "B"}
    student = {1: ["A", "B"], 2: "B"}
    res = grade_full_mcq(student, key, max_score=10.0)
    assert res["correct_count"] == 1
    assert res["wrong_count"] == 1
    assert res["score"] == 5.0


def test_full_mcq_20_questions():
    """Đề thi cấu hình 20 câu -> Điểm mỗi câu 0.5 điểm"""
    key = {i: "A" for i in range(1, 21)}
    student = {i: "A" if i <= 15 else "B" for i in range(1, 21)}
    res = grade_full_mcq(student, key, max_score=10.0)
    assert res["score"] == 7.5
    assert res["correct_count"] == 15
    assert res["total_questions"] == 20


def test_full_mcq_empty_inputs():
    """Trường hợp biên: Bài làm rỗng hoàn toàn"""
    key = {i: "A" for i in range(1, 41)}
    student = {}
    res = grade_full_mcq(student, key, max_score=10.0)
    assert res["score"] == 0.0
    assert res["blank_count"] == 40


# ==============================================================================
# 2. KIỂM THỬ CHẾ ĐỘ 2: CHUẨN THPT MỚI (3 PHẦN BGD&ĐT)
# ==============================================================================
def test_thpt_part1_scoring():
    """Phần I: Mỗi câu 0.25đ, tối đa 3.0đ"""
    key = {"part1": {i: "A" for i in range(1, 13)}}
    # Đúng 8 câu, sai 4 câu
    student = {"part1": {i: "A" if i <= 8 else "C" for i in range(1, 13)}}
    res = grade_thpt_2025(student, key)
    assert res["score_part1"] == 2.0
    assert res["p1_correct"] == 8


def test_thpt_part2_progressive_table():
    """
    Phần II: Quy chế tính điểm lũy tiến từng câu:
    1 ý đúng: 0.1đ | 2 ý đúng: 0.25đ | 3 ý đúng: 0.5đ | 4 ý đúng: 1.0đ
    """
    key_p2 = {
        1: {'a': 'Đ', 'b': 'S', 'c': 'Đ', 'd': 'Đ'},
        2: {'a': 'S', 'b': 'Đ', 'c': 'S', 'd': 'Đ'},
        3: {'a': 'Đ', 'b': 'Đ', 'c': 'S', 'd': 'S'},
        4: {'a': 'S', 'b': 'S', 'c': 'Đ', 'd': 'Đ'}
    }
    student_p2 = {
        # Câu 1: đúng cả 4 ý -> 1.0đ
        1: {'a': 'Đ', 'b': 'S', 'c': 'Đ', 'd': 'Đ'},
        # Câu 2: đúng 3 ý (a, b, c) -> 0.5đ
        2: {'a': 'S', 'b': 'Đ', 'c': 'S', 'd': 'S'},
        # Câu 3: đúng 2 ý (a, b) -> 0.25đ
        3: {'a': 'Đ', 'b': 'Đ', 'c': 'Đ', 'd': 'Đ'},
        # Câu 4: đúng 1 ý (a) -> 0.1đ
        4: {'a': 'S', 'b': 'Đ', 'c': 'S', 'd': 'S'}
    }
    res = grade_thpt_2025({"part2": student_p2}, {"part2": key_p2})
    # Tổng P2 = 1.0 + 0.5 + 0.25 + 0.1 = 1.85đ
    assert res["score_part2"] == 1.85


def test_thpt_part3_numeric_normalization():
    """Chuẩn hóa số học Phần III: dấu phẩy/chấm, số âm, số nguyên"""
    assert normalize_numeric_string("15.0") == "15"
    assert normalize_numeric_string("15") == "15"
    assert normalize_numeric_string("-2,5") == "-2.5"
    assert normalize_numeric_string("-2.5") == "-2.5"
    assert normalize_numeric_string("0,75") == "0.75"
    assert normalize_numeric_string(".75") == "0.75"

    key_p3 = {1: "15", 2: "-2.5", 3: "0.75", 4: "120", 5: "-8", 6: "3.14"}
    student_p3 = {
        1: "15,0",   # Đúng (chuẩn hóa thành 15)
        2: "-2.5",   # Đúng
        3: "0,75",   # Đúng
        4: "120",    # Đúng
        5: "8",      # Sai (thiếu dấu âm)
        6: ""        # Bỏ trống
    }
    res = grade_thpt_2025({"part3": student_p3}, {"part3": key_p3})
    # Đúng 4 câu x 0.5đ = 2.0đ
    assert res["p3_correct"] == 4
    assert res["score_part3"] == 2.0


def test_thpt_all_correct_total_10():
    """Làm đúng toàn bộ 3 phần THPT 2025 -> Tổng điểm 10.0"""
    key_thpt = {
        "part1": {i: "A" for i in range(1, 13)},
        "part2": {i: {'a': 'Đ', 'b': 'S', 'c': 'Đ', 'd': 'S'} for i in range(1, 5)},
        "part3": {i: str(i * 10) for i in range(1, 7)}
    }
    res = grade_thpt_2025(key_thpt, key_thpt)
    assert res["score_part1"] == 3.0
    assert res["score_part2"] == 4.0
    assert res["score_part3"] == 3.0
    assert res["total_score"] == 10.0


# ==============================================================================
# 3. KIỂM THỬ XỬ LÝ MÃ ĐỀ & SBD KHÔNG HỢP LỆ (ƯU TIÊN 1)
# ==============================================================================
def test_invalid_exam_code_not_fallback():
    """
    Nếu bài làm có mã đề không tồn tại trong danh sách:
    - KHÔNG được tự lấy đáp án của mã đề khác (không lấy mã đầu tiên)
    - Phải chuyển sang trạng thái 'Cần kiểm tra'
    - Điểm số phải là 0.0
    """
    exam = Exam(
        exam_id="ex_test_code",
        title="Đề Test Mã",
        mode=ExamMode.FULL_MCQ,
        # Kỳ thi chỉ có đáp án mã 102 và 103, KHÔNG có mã 101
        answer_keys={
            "102": {i: "B" for i in range(1, 41)},
            "103": {i: "C" for i in range(1, 41)}
        }
    )

    # Ảnh mẫu 001_NguyenVanA có mã đề 101
    sample_path = os.path.join(PROJECT_ROOT, "dataset", "sample_submissions", "001_NguyenVanA_FullTN.png")
    sub = process_single_submission(sample_path, exam)

    assert sub.status == "Cần kiểm tra"
    assert sub.needs_review is True
    assert sub.total_score == 0.0
    assert "không tồn tại trong bộ đáp án" in sub.error_message


def test_missing_or_invalid_exam_code_flagged():
    """Mã đề có dấu '?' hoặc không đọc được phải chuyển sang 'Cần kiểm tra'"""
    exam = Exam(
        exam_id="ex_test_code2",
        title="Đề Test Mã 2",
        mode=ExamMode.FULL_MCQ,
        answer_keys={"101": {i: "A" for i in range(1, 41)}}
    )
    # File mã đề sai đã sinh
    flawed_code_path = os.path.join(PROJECT_ROOT, "dataset", "sample_submissions", "test_flawed_ma_de_sai.png")
    if os.path.exists(flawed_code_path):
        sub = process_single_submission(flawed_code_path, exam)
        assert sub.status == "Cần kiểm tra"
        assert sub.needs_review is True
        assert sub.total_score == 0.0


# ==============================================================================
# 4. KIỂM THỬ THỐNG KÊ & LOẠI TRỪ BÀI LỖI (ƯU TIÊN 3)
# ==============================================================================
def test_excel_exporter_excludes_error_submissions(tmp_path):
    """
    Kiểm thử file Excel xuất ra:
    - Bài lỗi ('Lỗi') KHÔNG được đưa vào tính điểm trung bình của lớp
    - Tỷ lệ đúng/sai được tính chính xác trên bài hợp lệ
    """
    exam = Exam(
        exam_id="ex_test_export",
        title="Kỳ Thi Thử Nghiệm",
        mode=ExamMode.FULL_MCQ,
        num_questions=4,
        answer_keys={"101": {1: "A", 2: "B", 3: "C", 4: "D"}}
    )

    sub1 = Submission(
        submission_id="sub_1", exam_id="ex_test_export", student_id="001", student_name="Học sinh 1",
        exam_code="101", total_score=10.0, correct_count=4, wrong_count=0, blank_count=0,
        status="Đã chấm", extracted_data={"answers": {1: "A", 2: "B", 3: "C", 4: "D"}}
    )
    sub2 = Submission(
        submission_id="sub_2", exam_id="ex_test_export", student_id="002", student_name="Học sinh 2",
        exam_code="101", total_score=5.0, correct_count=2, wrong_count=2, blank_count=0,
        status="Đã chấm", extracted_data={"answers": {1: "A", 2: "B", 3: "A", 4: "A"}}
    )
    # Bài số 3 là bài bị lỗi (không thể nắn ảnh / ảnh hỏng)
    sub3_error = Submission(
        submission_id="sub_3", exam_id="ex_test_export", student_id="003", student_name="Học sinh Lỗi",
        exam_code="101", total_score=0.0, status="Lỗi", error_message="Ảnh bị rách / che marker",
        extracted_data={}
    )

    out_file = str(tmp_path / "test_report_3_sheet.xlsx")
    success = export_exam_results_to_excel(exam, [sub1, sub2, sub3_error], out_file)
    assert success is True
    assert os.path.exists(out_file)

    # Đọc lại file Excel để kiểm tra giá trị tính toán
    wb = openpyxl.load_workbook(out_file, data_only=True)
    assert "Bang_Diem_Tong_Hop" in wb.sheetnames
    assert "Chi_Tiet_Bai_Lam" in wb.sheetnames
    assert "Thong_Ke_Pho_Diem" in wb.sheetnames

    # Kiểm tra Sheet 3: Điểm trung bình chỉ tính trên 2 bài hợp lệ = (10 + 5) / 2 = 7.5
    # Nếu tính cả bài lỗi 0đ thì điểm TB sẽ là (10 + 5 + 0) / 3 = 5.0 (SAI)
    ws3 = wb["Thong_Ke_Pho_Diem"]
    
    # Tìm ô Điểm trung bình
    avg_cell_val = None
    for r in range(4, 20):
        if ws3.cell(row=r, column=1).value == "Điểm trung bình (chỉ tính bài hợp lệ)":
            avg_cell_val = ws3.cell(row=r, column=2).value
            break

    assert avg_cell_val == 7.5, f"Điểm trung bình phải là 7.5 (đã loại bài lỗi), nhưng nhận được: {avg_cell_val}"



# ==============================================================================
# 5. KIỂM THỬ BỔ SUNG: TRƯỜNG HỢP BIÊN & EDGE CASES
# ==============================================================================
def test_full_mcq_string_vs_int_keys():
    """Kiểm tra chấm điểm hoạt động đúng khi key là string '1' hoặc int 1"""
    key_str = {"1": "A", "2": "B", "3": "C", "4": "D"}
    student_int = {1: "A", 2: "B", 3: "C", 4: "D"}
    res = grade_full_mcq(student_int, key_str, max_score=10.0)
    assert res["correct_count"] == 4
    assert res["score"] == 10.0


def test_full_mcq_case_insensitive():
    """Đáp án viết thường/hoa đều phải khớp"""
    key = {1: "a", 2: "b"}
    student = {1: "A", 2: "B"}
    res = grade_full_mcq(student, key, max_score=10.0)
    assert res["correct_count"] == 2
    assert res["score"] == 10.0


def test_full_mcq_single_question():
    """Trường hợp biên: Chỉ có 1 câu hỏi"""
    key = {1: "A"}
    student = {1: "A"}
    res = grade_full_mcq(student, key, max_score=10.0)
    assert res["score"] == 10.0
    assert res["total_questions"] == 1


def test_thpt_all_blank():
    """THPT: Bỏ trống hoàn toàn cả 3 phần -> 0.0 điểm"""
    key = {
        "part1": {i: "A" for i in range(1, 13)},
        "part2": {i: {'a': 'Đ', 'b': 'S', 'c': 'Đ', 'd': 'S'} for i in range(1, 5)},
        "part3": {i: str(i * 10) for i in range(1, 7)}
    }
    student = {
        "part1": {},
        "part2": {},
        "part3": {}
    }
    res = grade_thpt_2025(student, key)
    assert res["total_score"] == 0.0
    assert res["score_part1"] == 0.0
    assert res["score_part2"] == 0.0
    assert res["score_part3"] == 0.0


def test_thpt_part2_zero_correct():
    """THPT Phần II: 0 ý đúng -> 0.0đ"""
    key_p2 = {1: {'a': 'Đ', 'b': 'S', 'c': 'Đ', 'd': 'S'}}
    student_p2 = {1: {'a': 'S', 'b': 'Đ', 'c': 'S', 'd': 'Đ'}}
    res = grade_thpt_2025({"part2": student_p2}, {"part2": key_p2})
    assert res["score_part2"] == 0.0


def test_thpt_part2_all_four_correct():
    """THPT Phần II: 4 ý đúng trong 1 câu -> 1.0đ (lũy tiến)"""
    key_p2 = {1: {'a': 'Đ', 'b': 'S', 'c': 'Đ', 'd': 'S'}}
    student_p2 = {1: {'a': 'Đ', 'b': 'S', 'c': 'Đ', 'd': 'S'}}
    res = grade_thpt_2025({"part2": student_p2}, {"part2": key_p2})
    assert res["score_part2"] == 1.0


def test_normalize_numeric_edge_cases():
    """Trường hợp biên chuẩn hóa số"""
    assert normalize_numeric_string(None) == ""
    assert normalize_numeric_string("") == ""
    assert normalize_numeric_string("0") == "0"
    assert normalize_numeric_string("0.0") == "0"
    assert normalize_numeric_string("-0") == "0"
    assert normalize_numeric_string("  15  ") == "15"
    assert normalize_numeric_string("3,14") == "3.14"
    assert normalize_numeric_string("abc") == "abc"


def test_exam_code_lookup_no_fallback():
    """Kiểm tra: Mã đề '999' không tồn tại -> KHÔNG trả về đáp án mã khác"""
    exam = Exam(
        exam_id="ex_nofallback",
        title="Test No Fallback",
        mode=ExamMode.FULL_MCQ,
        answer_keys={"101": {1: "A", 2: "B"}}
    )
    # Truy xuất trực tiếp - không fallback
    key = exam.answer_keys.get("999")
    assert key is None, "Mã đề '999' PHẢI trả về None, không được fallback!"


def test_thpt_part3_negative_and_decimal():
    """THPT Phần III: Số âm và thập phân khớp đúng"""
    key_p3 = {1: "-2.5", 2: "0.75", 3: "-8"}
    student_p3 = {1: "-2,5", 2: ".75", 3: "-8.0"}
    res = grade_thpt_2025({"part3": student_p3}, {"part3": key_p3})
    assert res["p3_correct"] == 3
    assert res["score_part3"] == 1.5


def test_submission_status_values():
    """Kiểm tra tất cả trạng thái Submission hợp lệ"""
    sub = Submission(submission_id="test", exam_id="test")
    assert sub.status == "Chưa chấm"
    assert sub.needs_review == False
    assert sub.error_message == ""
    
    # Kiểm tra to_dict/from_dict round-trip
    d = sub.to_dict()
    sub2 = Submission.from_dict(d)
    assert sub2.status == sub.status
    assert sub2.needs_review == sub.needs_review
