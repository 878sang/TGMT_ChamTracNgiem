"""
Bộ kiểm thử tự động cho Phần Mềm Nhận Diện & Chấm Phiếu Trắc Nghiệm THPT (50 Câu)
Chạy bằng: python tests/test_omr_50_cau.py
Sử dụng thư viện chuẩn unittest (không phụ thuộc pytest).
Kiểm tra toàn diện:
  - Chức năng 1: Quản lý và kiểm tra tính hợp lệ file ảnh
  - Chức năng 2: Tiền xử lý ảnh & Nắn phẳng phối cảnh (Perspective Transform)
  - Chức năng 3: Tự động dò ma trận 50 câu & bóc tách OMR (pixel density, bỏ trống, tô trùng)
  - Chức năng 4: Xử lý nhiều học sinh theo lô (Batch Processing & chống crash)
  - Chức năng 5: Chấm điểm tự động thang 10.0 & Tự dò đáp án chuẩn từ ảnh & Xuất Excel 3 Sheet
"""

import os
import sys
import unittest
import glob
import numpy as np

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

from core.models import Submission, Exam
from core.cv_utils import imread_utf8, imwrite_utf8
from core.image_alignment import align_paper, find_fiducial_markers, STANDARD_WIDTH, STANDARD_HEIGHT
from core.omr_extractor import extract_all_data, get_question_coords, detect_bubbles_50_questions
from core.grading_engine import grade_submission, parse_answer_string, generate_sample_answer_key
from core.exam_processor import process_single_submission, scan_answer_key_from_image
from core.excel_exporter import export_exam_results_to_excel
from core.sheet_generator import generate_all_samples, generate_50_mcq_template


class TestOMR50Cau(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        generate_all_samples()
        cls.answer_key = generate_sample_answer_key(50)
        cls.sample_dir = os.path.join(PROJECT_ROOT, "dataset", "sample_submissions_50cau")
        cls.template_dir = os.path.join(PROJECT_ROOT, "templates")

    # =========================================================================
    # CHỨC NĂNG 1: QUẢN LÝ VÀ KIỂM TRA ẢNH HỢP LỆ
    # =========================================================================
    def test_func1_read_valid_image(self):
        valid_path = os.path.join(self.sample_dir, "HS001_NguyenVanA.png")
        img = imread_utf8(valid_path)
        self.assertIsNotNone(img)
        self.assertEqual(img.shape[0], STANDARD_HEIGHT)
        self.assertEqual(img.shape[1], STANDARD_WIDTH)

    def test_func1_read_non_existent_image(self):
        img = imread_utf8("dataset/non_existent_image_12345.png")
        self.assertIsNone(img)

    # =========================================================================
    # CHỨC NĂNG 2: TIỀN XỬ LÝ & NẮN PHẲNG PHỐI CẢNH (PERSPECTIVE TRANSFORM)
    # =========================================================================
    def test_func2_align_flat_image(self):
        img_path = os.path.join(self.sample_dir, "HS001_NguyenVanA.png")
        success, aligned, msg = align_paper(img_path)
        self.assertTrue(success, msg)
        self.assertIsNotNone(aligned)
        self.assertEqual(aligned.shape[:2], (STANDARD_HEIGHT, STANDARD_WIDTH))

    def test_func2_align_skewed_camera_photo(self):
        img_path = os.path.join(self.sample_dir, "HS005_HoangVanE_ChupNghieng.jpg")
        success, aligned, msg = align_paper(img_path)
        self.assertTrue(success, f"Nắn phẳng ảnh chụp nghiêng thất bại: {msg}")
        self.assertEqual(aligned.shape[:2], (STANDARD_HEIGHT, STANDARD_WIDTH))

    def test_func2_align_blank_invalid_image(self):
        blank = np.ones((800, 600, 3), dtype=np.uint8) * 255
        success, aligned, msg = align_paper(blank)
        self.assertFalse(success)
        self.assertIsNone(aligned)

    # =========================================================================
    # CHỨC NĂNG 3: TỰ ĐỘNG DÒ MA TRẬN VÀ NHẬN DIỆN ĐÁP ÁN 50 CÂU
    # =========================================================================
    def test_func3_auto_detect_bubbles_50_questions(self):
        """Kiểm tra thuật toán OpenCV tự động dò ma trận 50 câu hỏi trên ảnh phẳng"""
        img_path = os.path.join(self.sample_dir, "HS001_NguyenVanA.png")
        img = imread_utf8(img_path, 0)
        coords_map = detect_bubbles_50_questions(img)
        self.assertEqual(len(coords_map), 50, "Thuật toán phải tự động dò ra đúng 50 câu hỏi")
        for q in range(1, 51):
            cy, cxs = coords_map[q]
            self.assertTrue(500 <= cy <= 1650)
            self.assertEqual(len(cxs), 4)

    def test_func3_detect_blank_questions(self):
        img_path = os.path.join(self.sample_dir, "HS003_LeVanC_BoTrong.png")
        sub = process_single_submission(img_path, self.answer_key)
        self.assertEqual(sub.blank_count, 6)
        for q in [3, 8, 15, 27, 33, 44]:
            self.assertIsNone(sub.answers.get(q))

    def test_func3_detect_multiple_marked_questions(self):
        img_path = os.path.join(self.sample_dir, "HS004_PhamThiD_ToTrung.png")
        sub = process_single_submission(img_path, self.answer_key)
        self.assertEqual(sub.multiple_count, 3)
        for q in [10, 22, 39]:
            self.assertIsInstance(sub.answers.get(q), list)

    # =========================================================================
    # CHỨC NĂNG 4: XỬ LÝ THEO LÔ (BATCH PROCESSING)
    # =========================================================================
    def test_func4_batch_processing_and_error_handling(self):
        all_images = sorted(glob.glob(os.path.join(self.sample_dir, "*.png")) + 
                            glob.glob(os.path.join(self.sample_dir, "*.jpg")))
        self.assertGreaterEqual(len(all_images), 7)

        results = []
        for p in all_images:
            sub = process_single_submission(p, self.answer_key)
            results.append(sub)

        self.assertEqual(len(results), len(all_images))
        success_list = [s for s in results if s.status == "Thành công"]
        self.assertGreater(len(success_list), 0)

    # =========================================================================
    # CHỨC NĂNG 5: CHẤM ĐIỂM & TỰ DÒ ĐÁP ÁN TỪ ẢNH & XUẤT EXCEL
    # =========================================================================
    def test_func5_grading_logic(self):
        st_ans = dict(self.answer_key)
        res = grade_submission(st_ans, self.answer_key, total_questions=50)
        self.assertEqual(res["score"], 10.0)
        self.assertEqual(res["correct_count"], 50)
        self.assertEqual(res["wrong_count"], 0)
        self.assertEqual(res["rank"], "Giỏi")

        for q in range(1, 11):
            st_ans[q] = 'B' if self.answer_key[q] != 'B' else 'C'
        res2 = grade_submission(st_ans, self.answer_key, total_questions=50)
        self.assertEqual(res2["score"], 8.0)
        self.assertEqual(res2["correct_count"], 40)
        self.assertEqual(res2["wrong_count"], 10)

    def test_func5_scan_answer_key_from_teacher_photo(self):
        """Kiểm tra tính năng tự động dò tìm đáp án chuẩn từ ảnh phiếu mẫu của giáo viên"""
        teacher_photo = os.path.join(self.template_dir, "phieu_dap_an_mau_giao_vien.png")
        success, scanned_key, msg = scan_answer_key_from_image(teacher_photo)
        self.assertTrue(success, f"Tự dò đáp án từ ảnh giáo viên thất bại: {msg}")
        self.assertEqual(len(scanned_key), 50, "Phải tự dò được đầy đủ 50 câu")
        # Kiểm tra xem đáp án dò được có khớp với key ban đầu không
        self.assertEqual(scanned_key[1], self.answer_key[1])
        self.assertEqual(scanned_key[50], self.answer_key[50])

    def test_func5_export_excel_3_sheets(self):
        import openpyxl
        out_excel = os.path.join(PROJECT_ROOT, "data", "test_report_50cau.xlsx")
        
        sub1 = Submission(submission_id="sub1", student_id="HS001", answers=self.answer_key, score=10.0, correct_count=50, status="Thành công")
        sub2 = Submission(submission_id="sub2", student_id="HS002", answers=self.answer_key, score=8.0, correct_count=40, wrong_count=10, status="Thành công")
        
        ok = export_exam_results_to_excel([sub1, sub2], self.answer_key, out_excel)
        self.assertTrue(ok)
        self.assertTrue(os.path.exists(out_excel))

        wb = openpyxl.load_workbook(out_excel)
        self.assertIn("Bang_Diem_Tong_Hop", wb.sheetnames)
        self.assertIn("Chi_Tiet_Bai_Lam", wb.sheetnames)
        self.assertIn("Thong_Ke_Pho_Diem", wb.sheetnames)


if __name__ == "__main__":
    suite = unittest.TestLoader().loadTestsFromTestCase(TestOMR50Cau)
    runner = unittest.TextTestRunner(verbosity=2)
    result = runner.run(suite)
    sys.exit(0 if result.wasSuccessful() else 1)
